"""기능사양서 로딩과 원문 위치 관리.

사양서는 라인 번호를 붙여 LLM에 전달하고, LLM이 돌려준 요구사항의 원문 인용(quote)과
라인 범위는 evidence.py에서 이 원문과 대조해 검증한다.
지원 형식: .txt .md .csv .rst (UTF-8/CP949), .pdf (pypdf 또는 pdftotext 필요), .docx (python-docx 필요)
"""

from __future__ import annotations

import re
import shutil
import subprocess
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

# 흔한 요구사항 ID 형식: SWR-BLV-001, REQ_123, [SRS-12], FR-1.2.3, SYS_REQ_0010
REQ_ID = re.compile(r"\b(?:[A-Z][A-Z0-9]{1,9}[-_])+(?:[A-Z]{0,4}[-_])?\d+(?:[._-]\d+)*\b")


@dataclass
class SpecDoc:
    name: str
    lines: list
    ids: list = field(default_factory=list)   # (id, line)

    def numbered(self) -> str:
        w = len(str(len(self.lines)))
        return "\n".join(f"S{n:0{w}d}| {t}" for n, t in enumerate(self.lines, 1))

    def text(self, start: int, end: int) -> str:
        return "\n".join(self.lines[max(0, start - 1): max(0, end)])

    def find_quote(self, quote: str, start: int | None = None, end: int | None = None, slack: int = 3):
        """인용문이 원문(정규화 비교)에 있는지 찾아 (시작, 끝) 라인을 돌려준다."""
        q = normalize(quote)
        if not q:
            return None
        n = len(self.lines)
        ranges = []
        if start:
            ranges.append((max(1, start - slack), min(n, (end or start) + slack)))
        ranges.append((1, n))
        for lo, hi in ranges:
            for i in range(lo, hi + 1):
                acc = ""
                for j in range(i, min(hi, i + 40) + 1):
                    acc += normalize(self.lines[j - 1])
                    if q in acc:
                        # 인용문을 포함하는 가장 늦은 시작 라인으로 좁힌다
                        k = i
                        while k < j and q in "".join(normalize(x) for x in self.lines[k: j]):
                            k += 1
                        return (k, j)
                    if len(acc) > len(q) * 4 + 200:
                        break
        return None


def normalize(s: str) -> str:
    s = unicodedata.normalize("NFKC", s or "")
    s = re.sub(r"[\s|*`>#_\-–—·•]+", "", s)
    return s.replace("“", "\"").replace("”", "\"").replace("‘", "'").replace("’", "'").lower()


def load_spec(path: str) -> SpecDoc:
    p = Path(path)
    suf = p.suffix.lower()
    if suf == ".pdf":
        text = _pdf_text(p)
    elif suf == ".docx":
        text = _docx_text(p)
    else:
        raw = p.read_bytes()
        for enc in ("utf-8-sig", "cp949", "euc-kr", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    while lines and not lines[-1].strip():
        lines.pop()
    doc = SpecDoc(p.name, lines)
    seen = set()
    for n, line in enumerate(lines, 1):
        for m in REQ_ID.finditer(line):
            if m.group(0) not in seen:
                seen.add(m.group(0))
                doc.ids.append((m.group(0), n))
    return doc


def _pdf_text(p: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore

        return "\n".join((page.extract_text() or "") for page in PdfReader(str(p)).pages)
    except ImportError:
        pass
    exe = shutil.which("pdftotext")
    if exe:
        r = subprocess.run([exe, "-layout", str(p), "-"], capture_output=True, text=True, errors="replace")
        if r.returncode == 0:
            return r.stdout
    raise ValueError(f"{p}: PDF 텍스트 추출에 pypdf 또는 pdftotext가 필요합니다 (pip install pypdf). "
                     "원문 라인 추적을 위해 텍스트/Markdown 변환본 사용을 권장합니다.")


def _docx_text(p: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as e:
        raise ValueError(f"{p}: .docx 읽기에 python-docx가 필요합니다 (pip install python-docx).") from e
    d = docx.Document(str(p))
    out = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)
